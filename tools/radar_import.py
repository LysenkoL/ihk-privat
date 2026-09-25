"""tools/radar_import.py — Themenanalyse (xlsx) → gen/radar-daten.js

Aufruf:  python tools/radar_import.py IHK_Pruefung_Themen.xlsx
Liest die Blätter „Teil 1“ und „Teil 2 FIAE“. Neue Themen müssen in META
(Anzeige-Name, Gruppe, Suchmuster) eingetragen werden, sonst bricht das Skript ab.
"""
import openpyxl, json, os, sys
HIER = os.path.dirname(os.path.abspath(__file__))
wb = openpyxl.load_workbook(sys.argv[1] if len(sys.argv) > 1 else "IHK_Pruefung_Themen.xlsx")

def mark(v):
    if v is None: return '.'
    s = str(v).strip().replace('\xa0', '')
    if not s: return '.'
    return '2' if s == '2' else 'x'

# label in xlsx -> (key, anzeige, gruppe, such-regex or None, extra)
META = {
 'Angebotskalk.': ('angebotskalk', 'Angebotskalkulation', 'kaufm', r'angebotskalkulation|listenverkaufspreis|barverkaufspreis|zielverkaufspreis', {}),
 'Handleskalk.': ('handelskalk', 'Handelskalkulation', 'kaufm', r'handelskalkulation|bezugspreis|listeneinkaufspreis|handelsspanne', {}),
 'Industriekalk.': ('industriekalk', 'Industriekalkulation', 'kaufm', r'industriekalkulation|herstellkosten|materialgemeinkosten|fertigungsgemeinkosten', {}),
 'Kostenkalkulation': ('kosten', 'Kostenkalkulation', 'kaufm', r'kostenvergleich|laufende[nr]? kosten|gesamtkosten|kosten (pro|je) (monat|jahr|arbeitsplatz|stück)|anschaffungskosten|kosten.{0,40}berechnen|berechnen.{0,60}kosten|stückkosten|fixkosten|variable kosten|kostenvergleichsrechnung', {}),
 'BAB': ('bab', 'Betriebsabrechnungsbogen', 'kaufm', r'betriebsabrechnungsbogen|\bbab\b', {}),
 'Deckungsbeitragsr.': ('deckungsbeitrag', 'Deckungsbeitragsrechnung', 'kaufm', r'deckungsbeitrag|break.?even|gewinnschwelle', {}),
 'Kalk. Wagniskosten': ('wagnis', 'Kalkulatorische Wagniskosten', 'kaufm', r'wagnis', {}),
 'Finanzierung/Leasing-Kauf': ('leasing', 'Finanzierung / Leasing oder Kauf', 'kaufm', r'leasing|finanzierung|kaufoption|ratenkauf|darlehen|kredit', {}),
 'Lagerkennzahlen': ('lager', 'Lagerkennzahlen', 'kaufm', r'lagerkennzahl|meldebestand|lagerumschlag|mindestbestand|lagerdauer', {}),
 'Wirtschaflichkeitsrechnung': ('wirtschaftlichkeit', 'Wirtschaftlichkeitsrechnung', 'kaufm', r'wirtschaftlichkeit|amortisation|rentabilit|return on invest', {}),
 'Buchhaltung': ('buchhaltung', 'Buchhaltung', 'kaufm', r'buchhaltung|buchungssatz|soll an haben', {}),
 'Aufbauorganisation': ('aufbauorg', 'Aufbauorganisation / Organigramm', 'kaufm', r'organigramm|aufbauorganisation|stabsstelle|linienorganisation|matrixorganisation', {}),
 'Kaufvertrag': ('kaufvertrag', 'Kaufvertrag & Störungen', 'kaufm', r'kaufvertrag|lieferverzug|lieferungsverzug|mangelhafte lieferung|annahmeverzug|zahlungsverzug|nicht.?rechtzeitig|mängelrüge|gewährleistung', {}),
 'Rechnungsbeleg': ('rechnung', 'Rechnung / Beleg prüfen', 'kaufm', r'\brechnung(en|sbeleg|spositionen|sprüfung|sstellung|sbetrag)?\b|\bbeleg|skonto|pflichtangaben|umsatzsteuer', {}),
 'Einkauf/Lager': ('einkauf', 'Einkauf / Lager', 'kaufm', r'einkauf|bestellmenge|bestellpunkt', {}),
 'Marketing': ('marketing', 'Marketing', 'kaufm', r'marketing|werbung|marktforschung|preispolitik', {}),
 'SWOT': ('swot', 'SWOT-Analyse', 'kaufm', r'swot', {}),
 'Angebotsvergleich': ('angebotsvergleich', 'Angebotsvergleich', 'kaufm', r'angebotsvergleich|angebote? .{0,30}vergleich|günstigste[sn]? angebot|rabatt', {}),
 'Marktformen': ('marktformen', 'Marktformen', 'kaufm', r'marktform|monopol|oligopol|polypol', {}),
 'Nutzwertanalyse': ('nutzwert', 'Nutzwertanalyse', 'projekt', r'nutzwert', {}),
 'Projektplanung': ('projekt', 'Projektplanung & -merkmale', 'projekt', r'projektmerkmal|merkmale eines projekt|projektplanung|projektphase|projektziel|magische[sn]? dreieck|projektauftrag|stakeholder|kick.?off', {}),
 'Lasten-/Pflichtenheft': ('lastenheft', 'Lasten- und Pflichtenheft', 'projekt', r'lastenheft|pflichtenheft', {}),
 'Projektstrukturplan': ('psp', 'Projektstrukturplan', 'projekt', r'projektstrukturplan|arbeitspaket', {}),
 'EPK': ('epk', 'EPK', 'projekt', r'\bepk\b|ereignisgesteuerte', {}),
 'Netzplan': ('netzplan', 'Netzplan', 'projekt', r'netzplan|kritische[rn]? pfad|gesamtpuffer|freie[rn]? puffer|\bfaz\b', {}),
 'Ganttdiagramm': ('gantt', 'Gantt-Diagramm', 'projekt', r'gantt|balkendiagramm', {}),
 'SMART': ('smart', 'SMART-Ziele', 'projekt', r'\bsmart\b', {}),
 'Datensicherheit': ('backup', 'Datensicherung (Backup)', 'sicher', r'datensicherung|backup|vollsicherung|inkrementell|differenziell|generationenprinzip|3-2-1|großvater', {}),
 'Verschlüsselung': ('verschluesselung', 'Verschlüsselung', 'sicher', r'verschlüssel|symmetrisch|asymmetrisch|öffentliche[nr]? schlüssel|private[nr]? schlüssel|public key|hybride', {}),
 'Elektronische Signatur': ('signatur', 'Elektronische Signatur & Zertifikate', 'sicher', r'signatur|zertifikat|hashwert|hashfunktion', {}),
 'Sicherheit/Schutz': ('schutz', 'IT-Sicherheit & Schutzziele', 'sicher', r'schutzziel|vertraulichkeit|integrität|verfügbarkeit|grundschutz|\bbsi\b|firewall|zugangskontrolle|zutrittskontrolle|zugriffskontrolle|passwort|authentifizierung|zwei-faktor|multi-faktor', {}),
 'Datenschutz': ('datenschutz', 'Datenschutz (DSGVO)', 'sicher', r'datenschutz|dsgvo|personenbezogen|\bbdsg\b|auftragsverarbeitung', {}),
 'Videokonferenzsysteme': ('video', 'Videokonferenzsysteme', 'system', r'videokonferenz', {}),
 'Support': ('support', 'Support & Service', 'system', r'support|ticket|helpdesk|service.?level|\bsla\b|first.?level|störungsmeldung', {}),
 'Berechnungen IT': ('speicher', 'Berechnungen: Speicher & Einheiten', 'system', r'speicher(bedarf|kapazität|platz)|\b(kib|mib|gib|tib|kibibyte|mebibyte|gibibyte|tebibyte)\b|farbtiefe|bildauflösung|auflösung von|farben .{0,20}darstellen', {}),
 'Energiekosten': ('energie', 'Energiekosten', 'system', r'energiekosten|stromkosten|\bkwh\b|stromverbrauch|energieverbrauch', {}),
 'Datenübertragung berechnen': ('uebertragung', 'Datenübertragung berechnen', 'system', r'übertragungs(dauer|zeit|rate)|download|upload|bandbreite|mbit/s|mbps|datenrate|speedtest', {}),
 'USV': ('usv', 'USV', 'system', r'\busv\b|unterbrechungsfreie', {}),
 'Übersetzen': ('englisch', 'Englisch / Übersetzen', 'sonst', r'übersetz|englisch|english|translate', {}),
 'IT-Struktur - Homeoffice': ('homeoffice', 'IT im Homeoffice', 'system', r'home.?office|telearbeit|mobiles arbeiten', {}),
 'Allgemeine Begriffe': ('begriffe', 'Allgemeine IT-Begriffe', 'sonst', None, {}),
 'Betriebssysteme': ('bs', 'Betriebssysteme', 'system', r'betriebssystem', {}),
 'OSI': ('osi', 'OSI-Modell', 'netz', r'\bosi\b|osi-', {}),
 'Datenspeicher RAID': ('raid', 'RAID', 'system', r'\braid\b', {'ap2': True}),
 'Programmierung': ('programmierung', 'Programmierung (Grundlagen)', 'software', r'programmier|quellcode|variable|konstante|compiler|interpreter|objektorient|prozedural|datentyp', {}),
 'Struktogramm': ('struktogramm', 'Struktogramm', 'software', r'struktogramm|nassi', {}),
 'Klassendiagramm + OOP': ('oop', 'Klassendiagramm & OOP', 'software', r'klassendiagramm|objektorient|vererbung|kapselung|\bklasse\b|\bobjekt(e)?\b', {}),
 'Algorithmus': ('algorithmus', 'Algorithmus / Pseudocode', 'software', r'algorithmus|pseudocode|quellcode|schleife|funktion .{0,40}\(', {}),
 'Schreibtischtest': ('schreibtisch', 'Schreibtischtest', 'software', r'schreibtischtest|führen sie .{0,30}test durch|rückgabe der funktion|rückgabewert', {}),
 'UML UseCase Diagramm': ('usecase', 'UML-Anwendungsfalldiagramm', 'software', r'anwendungsfall|use.?case', {}),
 'UML Aktivitätsdiagramm': ('aktivitaet', 'UML-Aktivitätsdiagramm', 'software', r'aktivitätsdiagramm', {}),
 'Webentwicklung': ('web', 'Webentwicklung', 'software', r'webseite|website|\bhtml\b|\bcss\b|javascript|webanwendung|responsive|\bseo\b', {}),
 'Datenbank': ('datenbank', 'Datenbank-Grundlagen', 'db', r'datenbank|fremdschlüssel|primärschlüssel|relational', {}),
 'ERM': ('erm', 'ER-Modell', 'db', r'entity|er-diagramm|er-modell|datenmodell|kardinalität|entität', {}),
 'Datenbank-Modell': ('relmodell', 'Relationales Modell', 'db', r'relationale[sn]? (daten)?modell|tabellenmodell|normalform|normalisier', {}),
 'SQL': ('sql', 'SQL', 'db', r'\bsql\b', {'ap2': True}),
 'IP v4 & v6': ('ip', 'IPv4 & IPv6 / Subnetting', 'netz', r'ipv4|ipv6|ip-adresse|subnetz|netzadresse|broadcast|präfix|\bcidr\b', {}),
 'Netzwerk': ('netzwerk', 'Netzwerk & Fehlersuche', 'netz', r'netzwerk|switch|router|\blan\b|\bping\b|ipconfig|\bdhcp\b|\bdns\b|gateway|mac-adresse|\barp\b|tracert|nslookup', {}),
 'Cloud': ('cloud', 'Cloud (IaaS/PaaS/SaaS)', 'netz', r'cloud|\bsaas\b|\biaas\b|\bpaas\b', {}),
 'VPN': ('vpn', 'VPN', 'netz', r'\bvpn\b|virtual private', {}),
 'File-Server': ('fileserver', 'File-Server / NAS', 'netz', r'file.?server|\bnas\b|dateiserver', {}),
 'WLAN': ('wlan', 'WLAN', 'netz', r'wlan|\bwpa|access.?point|\bssid\b', {}),
 'Virtualisierung': ('virtualisierung', 'Virtualisierung', 'system', r'virtualisier|hypervisor|virtuelle[nr]? maschine', {}),
 'Hardware': ('hardware', 'Hardware & Schnittstellen', 'system', r'hardware|prozessor|\bcpu\b|arbeitsspeicher|\bram\b|mainboard|netzteil|grafikkarte|schnittstelle|\bssd\b|\bhdd\b|\busb\b|displayport|hdmi|taktfrequenz', {}),
 'Hardware-Software ': ('hwsw', 'Hardware-/Software-Auswahl', 'system', r'hardware.{0,20}software|softwareauswahl', {}),
 'Malware, Phishing': ('malware', 'Malware & Phishing', 'sicher', r'malware|phishing|virus|viren|trojaner|ransomware|\bwurm|schadsoftware|social engineering', {}),
 'Lizenzmodelle': ('lizenz', 'Lizenzmodelle', 'kaufm', r'lizenz|open.?source|pay.?per.?use|subskription|abonnement|\beula\b|freeware', {}),
 'Einsatz von KI': ('ki', 'Einsatz von KI', 'software', r'\bki\b|künstliche[rn]? intelligenz|chatbot|machine learning|\bki-', {}),
 'P = U * I': ('strom', 'Leistung P = U · I', 'system', r'p ?= ?u|leistungsaufnahme|\bwatt\b|\bvolt\b|ampere|stromstärke|\bpoe\b|netzteil', {}),
 'Ergonomischer Arbeitsplatz': ('ergonomie', 'Ergonomischer Arbeitsplatz', 'system', r'ergonom|bildschirmarbeit|körperhaltung', {}),
 'Daisy-Chaining': ('daisy', 'Daisy-Chaining (Monitore)', 'system', r'daisy|\bmst\b|thunderbolt|displayport', {}),
}

ws = wb['Teil 1']
rows = list(ws.iter_rows(values_only=True))
teil1 = []
unbekannt = []
for r in rows[2:]:
    name = r[0]
    if not name or not str(name).strip(): continue
    alt = ''.join(mark(v) for v in r[1:22])
    neu = ''.join(mark(v) for v in r[22:32])
    if 'x' not in alt + neu and '2' not in alt + neu: continue
    if name not in META:
        unbekannt.append(name); continue
    key, anz, gr, such, extra = META[name]
    e = {'k': key, 't': anz, 'g': gr, 'alt': alt, 'neu': neu}
    if such: e['such'] = such
    e.update(extra)
    teil1.append(e)
if unbekannt:
    sys.exit('Unbekannte Themen, bitte in META eintragen: ' + ', '.join(map(str, unbekannt)))

ws = wb['Teil 2 FIAE']
rows = list(ws.iter_rows(values_only=True))
ap2 = []
for r in rows[2:]:
    name = r[0]
    if not name or not str(name).strip(): continue
    alt = ''.join(mark(v) for v in r[1:14])
    neu = ''.join(mark(v) for v in r[15:24])
    if 'x' not in alt + neu and '2' not in alt + neu: continue
    ap2.append({'t': str(name).strip(), 'alt': alt, 'neu': neu})

daten = {
  'stand': __import__('datetime').date.today().isoformat(),
  'quelle': 'Themenanalyse IHK_Pruefung_Themen.xlsx (Teil 1, Teil 2 FIAE)',
  'spalten': ['H21', 'F22', 'H22', 'F23', 'H23', 'F24', 'H24', 'F25', 'H25', 'F26'],
  'exams': ['ap1-2021-h', 'ap1-2022-f', 'ap1-2022-h', 'ap1-2023-f', 'ap1-2023-h', 'ap1-2024-f', 'ap1-2024-h', 'ap1-2025-f', 'ap1-2025-h', 'ap1-2026-f'],
  'altSpalten': 'Sommer 2013 bis Sommer 2023 (alte Ausbildungsordnung, Teil 1 = „GA1“), 21 Prüfungen',
  'ap2Spalten': ['W21/22', 'S22', 'W22/23', 'S23', 'W23/24', 'S24', 'W24/25', 'S25', 'W25/26'],
  'gruppen': {'kaufm': 'Kaufmännisch & Recht', 'projekt': 'Projektmanagement', 'sicher': 'IT-Sicherheit & Datenschutz',
              'system': 'IT-Systeme & Berechnungen', 'netz': 'Netzwerk', 'software': 'Software & Programmierung',
              'db': 'Datenbanken', 'sonst': 'Sonstiges'},
  'teil1': teil1,
  'ap2': ap2,
}
js = json.dumps(daten, ensure_ascii=False, indent=1)
kopf = '''/* ============================================================================
   gen/radar-daten.js — welches Thema kam in welcher Prüfung dran?
   ----------------------------------------------------------------------------
   Aus der Themenanalyse „IHK_Pruefung_Themen.xlsx“ (Blätter „Teil 1“ und
   „Teil 2 FIAE“). Je Thema zwei Zeichenketten, eine Stelle je Prüfung:
     x = Thema kam dran, . = nicht, 2 = in zwei Aufgaben (nur AP2)
   alt = alte Ausbildungsordnung (Sommer 2013 … Sommer 2023, 21 Prüfungen)
   neu = neue AP1 (Herbst 2021 … Frühjahr 2026, 10 Prüfungen) — dieselben zehn
         Prüfungen, die in der App liegen, Reihenfolge wie `spalten`/`exams`.
   Die Spalte „Anzahl“ der Tabelle ist nicht übernommen: sie zählte teilweise
   falsch (Malware 28 statt 2). Gezählt wird hier aus den Kreuzen selbst.
   such = Suchmuster (klein geschrieben), mit dem gen/radar.js passende
          Aufgaben aus den IHK-Prüfungen und dem Azubi-Navigator findet.
   ap2 = laut Prüfungskatalog nur noch in Teil 2 (SQL, RAID).
   Erzeugt mit tools/radar_import.py — bei neuen Daten neu erzeugen.
   ========================================================================== */
'''
open(os.path.join(HIER, '..', 'gen', 'radar-daten.js'), 'w', encoding='utf-8').write(kopf + '(function (root) {\n  const D = ' + js.replace('\n', '\n  ') + ';\n  root.IHK_RADAR = D;\n  if (typeof module === "object" && module.exports) module.exports = D;\n})(typeof window !== "undefined" ? window : globalThis);\n')
print('gen/radar-daten.js:', len(teil1), 'Themen Teil 1,', len(ap2), 'Themen AP2 FIAE')
